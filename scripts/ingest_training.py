"""Stream a local training Excel workbook into a persistent Chroma collection.

Run from the repository root:
    python -m scripts.ingest_training
"""
import hashlib
import os
import sys
from pathlib import Path
from typing import Iterator

import chromadb
from openpyxl import load_workbook

from app.services.embedding_service import get_embedding_function

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_EXCEL_PATH = PROJECT_ROOT / "data" / "training.xlsx"
DEFAULT_CHROMA_PATH = PROJECT_ROOT / "chroma_db"
COLLECTION_NAME = "training_catalogue"
BATCH_SIZE = int(os.getenv("TRAINING_INGEST_BATCH_SIZE", "128"))

COLUMN_MAP = {
    "Series Name": "series_name",
    "Stage": "stage",
    "Architecture": "architecture",
    "Program Track": "program_track",
    "Track": "track",
    "Role": "role",
    "Language": "language",
    "Learning Plan Type": "learning_plan_type",
    "ENABLEMENT_BUSINESS_ENTITY": "business_entity",
    "Description": "description",
    "Series URL": "series_url",
}
REQUIRED_COLUMNS = {"Series Name", "Description", "Series URL"}


def clean_value(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def stable_id(row: dict[str, str]) -> str:
    """Stable ID lets upsert update existing training when its metadata changes."""
    identity = row["series_url"].strip() or (
        row["series_name"].strip() + "|" + row["description"].strip()
    )
    return hashlib.sha256(identity.casefold().encode("utf-8")).hexdigest()


def document_text(row: dict[str, str]) -> str:
    fields = [
        ("Training", row["series_name"]),
        ("Description", row["description"]),
        ("Stage", row["stage"]),
        ("Architecture", row["architecture"]),
        ("Program track", row["program_track"]),
        ("Track", row["track"]),
        ("Role", row["role"]),
        ("Language", row["language"]),
        ("Learning plan type", row["learning_plan_type"]),
        ("Business entity", row["business_entity"]),
    ]
    return "\n".join(f"{label}: {value}" for label, value in fields if value)


def read_rows(excel_path: Path) -> Iterator[dict[str, str]]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = [clean_value(value) for value in next(rows, ())]
        header_indexes = {name: index for index, name in enumerate(headers) if name}

        missing = REQUIRED_COLUMNS - set(header_indexes)
        if missing:
            raise ValueError(
                "Excel is missing required columns: " + ", ".join(sorted(missing))
            )

        for row_number, values in enumerate(rows, start=2):
            raw = {}
            for excel_name in COLUMN_MAP:
                index = header_indexes.get(excel_name)
                raw[excel_name] = clean_value(values[index]) if index is not None and index < len(values) else ""

            if not raw["Series Name"] and not raw["Description"]:
                continue

            yield {COLUMN_MAP[key]: value for key, value in raw.items()}
    finally:
        workbook.close()


def main() -> int:
    excel_path = Path(os.getenv("TRAINING_EXCEL_PATH", str(DEFAULT_EXCEL_PATH)))
    chroma_path = Path(os.getenv("CHROMA_PERSIST_DIRECTORY", str(DEFAULT_CHROMA_PATH)))

    if not excel_path.is_file():
        print(
            f"Training Excel file not found: {excel_path}\n"
            "Place your workbook at data/training.xlsx, or set TRAINING_EXCEL_PATH.",
            file=sys.stderr,
        )
        return 1

    chroma_path.mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=str(chroma_path))
    collection = client.get_or_create_collection(
        name=COLLECTION_NAME,
        embedding_function=get_embedding_function(),
        metadata={"hnsw:space": "cosine"},
    )

    current_ids: set[str] = set()
    ids_batch: list[str] = []
    documents_batch: list[str] = []
    metadatas_batch: list[dict[str, str]] = []
    processed = 0

    def flush_batch():
        nonlocal processed
        if not ids_batch:
            return
        collection.upsert(
            ids=ids_batch,
            documents=documents_batch,
            metadatas=metadatas_batch,
        )
        processed += len(ids_batch)
        ids_batch.clear()
        documents_batch.clear()
        metadatas_batch.clear()
        print(f"Indexed/updated {processed} training records...")

    try:
        for row in read_rows(excel_path):
            training_id = stable_id(row)
            if training_id in current_ids:
                print(
                    "Warning: duplicate training identity skipped: "
                    + (row["series_name"] or row["series_url"])
                )
                continue

            current_ids.add(training_id)
            ids_batch.append(training_id)
            documents_batch.append(document_text(row))
            metadatas_batch.append(row)

            if len(ids_batch) >= BATCH_SIZE:
                flush_batch()

        flush_batch()
    except Exception as exc:
        print(f"Ingestion failed: {exc}", file=sys.stderr)
        return 1

    # Remove records no longer present in the workbook, but only after all batches succeeded.
    existing = collection.get(include=[])
    stale_ids = [item for item in existing["ids"] if item not in current_ids]
    for start in range(0, len(stale_ids), BATCH_SIZE):
        collection.delete(ids=stale_ids[start : start + BATCH_SIZE])

    print(
        f"Training ingestion complete. Workbook records: {len(current_ids)}; "
        f"collection count: {collection.count()}; removed stale records: {len(stale_ids)}."
    )
    print(f"Chroma data directory: {chroma_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
